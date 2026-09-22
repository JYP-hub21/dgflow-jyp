# -*- coding: utf-8 -*-
"""
발주서 원본에서 "입력값만" 비워 빈 틀을 만든다.

수식·서식·병합·열너비는 그대로 두고, 사람이 손으로 적은 값만 지운다.
그래서 결과 파일은 원본과 똑같이 생겼지만 비어 있고, 채우면 평수·실리콘·발주수량이 자동 계산된다.

사용법:
  python scripts/make-blank-form.py "<원본.xlsx>" "<시트이름>" "<결과.xlsx>"
"""
import sys
from openpyxl import load_workbook

# 지우지 않는 글자 — 양식의 칸 이름들
LABELS = {'계', '기본', '발주세대', '카자리', '발주수량', '위 치', '제 품 명', '규      격',
          '수 량', '평 수', '실리콘', '비 고', '두께(mm)', '가 로', '세 로'}

# 머리 영역에서 비울 칸 (값이 들어가는 자리만)
HEADER_INPUTS = ['H2', 'B4', 'H4', 'B5', 'G5', 'B6', 'G6', 'B7', 'G7', 'A11', 'A12']
DATA_START_ROW = 15


def main(src, sheet, dst):
    wb = load_workbook(src)                       # 수식을 살리려고 data_only=False
    for name in list(wb.sheetnames):              # 요청한 시트만 남긴다
        if name != sheet:
            del wb[name]
    ws = wb[sheet]
    ws.title = '발주서'

    cleared = 0
    for addr in HEADER_INPUTS:
        if ws[addr].value not in (None, ''):
            ws[addr].value = None
            cleared += 1

    for row in ws.iter_rows(min_row=DATA_START_ROW, max_row=ws.max_row):
        for cell in row:
            v = cell.value
            if v is None or v == '':
                continue
            if isinstance(v, str) and v.startswith('='):   # 수식은 남긴다
                continue
            if isinstance(v, str) and v.strip() in LABELS:  # 칸 이름은 남긴다
                continue
            cell.value = None
            cleared += 1

    wb.save(dst)
    print(f'비운 칸 {cleared}개 · 저장: {dst}')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2], sys.argv[3])
