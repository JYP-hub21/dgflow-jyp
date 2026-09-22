# -*- coding: utf-8 -*-
"""
발주서 → 규격정리 변환기

우미(=KCC) 양식 발주서를 읽어, 담당자가 손으로 만들던 "규격정리" 엑셀을 그대로 만든다.
규격정리 파일들(규격정리-우미건설.xlsx 등)의 배치를 따른다:

  A         B     C     D     E     F   G         H     I     J     K     L
  [묶음]    품명        날짜        로이=   품명
  위치      규    격    수량  평        위치      규    격    수량  평    M2
  (원본 순서 그대로)                    (합산·정렬한 정리본)
  ...                                   ...
            계 수량  평                  계 수량  평   M2

규칙 (실제 ERP 내역 9장으로 검증):
  1 합산    품명 + 가로 + 세로 같으면 수량을 더한다 (방이 달라도)
  2 정렬    세로 내림차순 → 가로 내림차순
  3 호 접기  3개 이상 연속이면 1~3호, 아니면 4,5호
  4 타입·방 접기  84A,B / 주방/식당,드레스룸
  5 검산    복층은 外·內가 짝 — 한쪽만 있으면 경고

묶음(작업의뢰서 한 장이 될 단위)은 사람이 정한다. 문장은 해석하지 않는다.

사용법:
  python scripts/spec-organize.py "<발주서.xlsx>" --sheet 7차 --bundles "1:1,2,3|4,5;2:1,2|3,4,5" [--out 결과.xlsx]
  --bundles 를 주지 않으면 동마다 한 묶음으로 본다.
"""
import argparse
import re
import sys
from collections import Counter, OrderedDict
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

JAE = 0.303 * 0.303          # 1才 = 0.091809 ㎡
DATA_START_ROW = 15


# ── 읽기 ────────────────────────────────────────────────────────────────
def num(v):
    try:
        return float(str(v).replace(',', '').strip() or 0)
    except ValueError:
        return 0.0


def read_order(path, sheet):
    """구조화된 열(I 동호 · J 실명 · K 타입 층)을 읽는다. C열 위치 문장은 원본 표시용으로만 둔다."""
    wb = load_workbook(path, data_only=True)      # 수식은 엑셀이 저장해 둔 결과값으로
    if sheet not in wb.sheetnames:
        sys.exit(f'시트 "{sheet}" 가 없습니다. 있는 시트: {wb.sheetnames}')
    ws = wb[sheet]
    items, product = [], ''
    for r in range(DATA_START_ROW, ws.max_row + 1):
        b = ws.cell(r, 2).value
        if b not in (None, ''):
            product = str(b).strip()              # 품명은 블록 첫 줄에만 있어 아래로 채운다
        dh = re.match(r'^(\d+)동 (\d+)호$', str(ws.cell(r, 9).value or '').strip())
        tk = re.match(r'^(\S+)\s+(\S+)$', str(ws.cell(r, 11).value or '').strip())
        room = str(ws.cell(r, 10).value or '').strip()
        w, h, q = num(ws.cell(r, 4).value), num(ws.cell(r, 5).value), num(ws.cell(r, 6).value)
        if not (dh and tk and room and w and h) or q == 0:
            continue
        items.append(dict(
            row=r, product=product, dong=int(dh[1]), ho=int(dh[2]), type=tk[1], floor=tk[2],
            room=room, w=int(w), h=int(h), qty=int(q),
            raw_loc=str(ws.cell(r, 3).value or '').strip(),
        ))
    header = dict(
        site=str(ws['B4'].value or '').strip(),
        order_no=str(ws['H2'].value or '').strip(),
        order_date=str(ws['B6'].value or '').strip(),
        note=' '.join(str(ws[a].value or '').strip() for a in ('A11', 'A12')).strip(),
    )
    return header, items


# ── 묶음 ────────────────────────────────────────────────────────────────
def parse_bundles(spec, items):
    """"1:1,2,3|4,5;2:1,2|3,4,5" → [{id, dong, ho:[...]}, ...]. 없으면 동마다 한 묶음."""
    bundles = []
    if spec:
        for part in spec.split(';'):
            dong, groups = part.split(':')
            for g in groups.split('|'):
                bundles.append(dict(dong=int(dong), ho=[int(x) for x in g.split(',') if x.strip()]))
    else:
        for dong in sorted({it['dong'] for it in items}):
            bundles.append(dict(dong=dong, ho=sorted({it['ho'] for it in items if it['dong'] == dong})))
    for i, b in enumerate(bundles, 1):
        b['id'] = f'G{i}'
    return bundles


def fold_nums(nums):
    s = sorted(set(nums))
    if len(s) == 1:
        return str(s[0])
    seq = len(s) >= 3 and all(b - a == 1 for a, b in zip(s, s[1:]))
    return f'{s[0]}~{s[-1]}' if seq else ','.join(map(str, s))


def fold_types(types):
    s = sorted(set(types))
    if len(s) == 1:
        return s[0]
    base = re.sub(r'[A-Z]+$', '', s[0])
    if all(t.startswith(base) for t in s):
        return base + ','.join(t[len(base):] for t in s)
    return ','.join(s)


def fold_rooms(rooms):
    m = re.search(r'\((外|內)\)$', rooms[0])
    suffix = m.group(0) if m else ''
    names = OrderedDict()
    for r in rooms:
        for n in re.sub(r'\((外|內)\)$', '', r).split(','):
            names[n] = True
    return ','.join(names) + suffix


def mode(values):
    return Counter(values).most_common(1)[0][0]


def organize(items, bundle):
    """묶음 하나를 규격정리 블록들로: 품명별 {left: 원본 순서, right: 합산·정렬}"""
    mine = [it for it in items if it['dong'] == bundle['dong'] and it['ho'] in bundle['ho']]
    blocks = OrderedDict()
    for it in mine:
        blocks.setdefault(it['product'], dict(left=[], groups=OrderedDict()))
        blk = blocks[it['product']]
        blk['left'].append(it)
        g = blk['groups'].setdefault((it['w'], it['h']), dict(
            w=it['w'], h=it['h'], qty=0, hos=[], types=[], rooms=[], floors=[], dong=it['dong']))
        g['qty'] += it['qty']; g['hos'].append(it['ho']); g['types'].append(it['type'])
        g['rooms'].append(it['room']); g['floors'].append(it['floor'])
    for blk in blocks.values():
        right = []
        for g in blk['groups'].values():
            loc = f"{g['dong']}동 {fold_nums(g['hos'])}호 {fold_types(g['types'])} {mode(g['floors'])} {fold_rooms(g['rooms'])}"
            right.append(dict(loc=loc, w=g['w'], h=g['h'], qty=g['qty']))
        right.sort(key=lambda x: (-x['h'], -x['w']))
        blk['right'] = right
    return mine, blocks


def check_pairs(mine):
    """外/內 짝 검산"""
    key = lambda it: (it['dong'], it['ho'], it['w'], it['h'], re.sub(r'\((外|內)\)$', '', it['room']))
    outer = {key(i): i for i in mine if i['room'].endswith('(外)')}
    inner = {key(i): i for i in mine if i['room'].endswith('(內)')}
    warns = []
    for k, it in outer.items():
        if k not in inner:
            warns.append(f"外에는 있는데 內가 없음 → {k[0]}동 {k[1]}호 {k[2]}x{k[3]} {k[4]}  {it['qty']}매")
    for k, it in inner.items():
        if k not in outer:
            warns.append(f"內에는 있는데 外가 없음 → {k[0]}동 {k[1]}호 {k[2]}x{k[3]} {k[4]}  {it['qty']}매")
    return warns


# ── 쓰기 ────────────────────────────────────────────────────────────────
THIN = Side(style='thin', color='999999')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill('solid', fgColor='EEEEEE')
BUNDLE_FILL = PatternFill('solid', fgColor='DCEBE3')
BOLD = Font(bold=True)
CENTER = Alignment(horizontal='center', vertical='center')


def area_jae(w, h, q):
    return round(w * h * q / 1_000_000 / JAE, 2)


def area_m2(w, h, q):
    return round(w * h * q / 1_000_000, 2)


def write_spec(path, header, bundles, results, warnings):
    wb = Workbook()
    ws = wb.active
    ws.title = '규격정리'
    widths = [44, 8, 8, 7, 9, 4, 44, 8, 8, 7, 9, 9]
    for i, wd in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    today = date.today().strftime('%y.%m.%d')

    r = 1
    for b, (mine, blocks) in zip(bundles, results):
        # 묶음 띠
        ws.cell(r, 1, f"{b['id']}  {b['dong']}동 {fold_nums(b['ho'])}호   ({len(mine)}줄 → {sum(len(x['right']) for x in blocks.values())}줄)")
        ws.cell(r, 1).font = BOLD
        for c in range(1, 13):
            ws.cell(r, c).fill = BUNDLE_FILL
        r += 1
        for product, blk in blocks.items():
            # 블록 머리: 품명 · 날짜 · 품명
            ws.cell(r, 2, product).font = BOLD
            ws.cell(r, 5, today)
            ws.cell(r, 8, product).font = BOLD
            r += 1
            for c, t in zip(range(1, 13), ['위치', '규', '격', '수량', '평', '', '위치', '규', '격', '수량', '평', 'M2']):
                cell = ws.cell(r, c, t)
                cell.fill = HEAD_FILL; cell.alignment = CENTER; cell.border = BOX; cell.font = BOLD
            r += 1
            top = r
            n = max(len(blk['left']), len(blk['right']))
            for i in range(n):
                if i < len(blk['left']):
                    it = blk['left'][i]
                    ws.cell(r, 1, it['raw_loc'] or f"{it['dong']}동 {it['ho']}호 {it['type']} {it['floor']} {it['room']}")
                    ws.cell(r, 2, it['w']); ws.cell(r, 3, it['h']); ws.cell(r, 4, it['qty'])
                    ws.cell(r, 5, area_jae(it['w'], it['h'], it['qty']))
                if i < len(blk['right']):
                    x = blk['right'][i]
                    ws.cell(r, 6, i + 1)
                    ws.cell(r, 7, x['loc']); ws.cell(r, 8, x['w']); ws.cell(r, 9, x['h']); ws.cell(r, 10, x['qty'])
                    ws.cell(r, 11, area_jae(x['w'], x['h'], x['qty']))
                    ws.cell(r, 12, area_m2(x['w'], x['h'], x['qty']))
                for c in range(1, 13):
                    ws.cell(r, c).border = BOX
                for c in (2, 3, 8, 9):
                    ws.cell(r, c).number_format = '#,##0'
                for c in (5, 11, 12):
                    ws.cell(r, c).number_format = '0.00'
                r += 1
            # 계
            ws.cell(r, 4, f'=SUM(D{top}:D{r-1})'); ws.cell(r, 5, f'=SUM(E{top}:E{r-1})')
            ws.cell(r, 10, f'=SUM(J{top}:J{r-1})'); ws.cell(r, 11, f'=SUM(K{top}:K{r-1})'); ws.cell(r, 12, f'=SUM(L{top}:L{r-1})')
            for c in (4, 5, 10, 11, 12):
                ws.cell(r, c).font = BOLD; ws.cell(r, c).border = BOX
                ws.cell(r, c).number_format = '0.00' if c in (5, 11, 12) else '#,##0'
            r += 2
        r += 1

    # 검산 시트
    ws2 = wb.create_sheet('검산')
    ws2.column_dimensions['A'].width = 14; ws2.column_dimensions['B'].width = 90
    ws2.append(['현장', header['site']]); ws2.append(['발주번호', header['order_no']])
    ws2.append(['발주일', header['order_date']]); ws2.append(['고객 지시 원문', header['note']])
    ws2.append([])
    ws2.append(['묶음', '동·호', '원본 줄', '정리 줄', '총 수량'])
    for b, (mine, blocks) in zip(bundles, results):
        ws2.append([b['id'], f"{b['dong']}동 {fold_nums(b['ho'])}호", len(mine),
                    sum(len(x['right']) for x in blocks.values()), sum(i['qty'] for i in mine)])
    ws2.append([])
    ws2.append(['검산 결과', '없음' if not warnings else f'{len(warnings)}건 — 아래 확인'])
    for w in warnings:
        ws2.append(['', w])
    wb.save(path)


# ── 실행 ────────────────────────────────────────────────────────────────
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('src'); ap.add_argument('--sheet', required=True)
    ap.add_argument('--bundles', default=''); ap.add_argument('--out', default='')
    a = ap.parse_args()

    header, items = read_order(a.src, a.sheet)
    bundles = parse_bundles(a.bundles, items)
    results, warnings = [], []
    for b in bundles:
        mine, blocks = organize(items, b)
        results.append((mine, blocks))
        warnings += [f"{b['id']}  {w}" for w in check_pairs(mine)]
    unassigned = [it for it in items if not any(it['dong'] == b['dong'] and it['ho'] in b['ho'] for b in bundles)]

    out = a.out or re.sub(r'\.xlsx$', '', a.src) + f'_규격정리({a.sheet}).xlsx'
    write_spec(out, header, bundles, results, warnings)

    print(f"읽음: {len(items)}줄 · 품명 {len({i['product'] for i in items})}종 · 묶음 {len(bundles)}개")
    for b, (mine, blocks) in zip(bundles, results):
        print(f"  {b['id']}  {b['dong']}동 {fold_nums(b['ho'])}호  {len(mine):3d}줄 → {sum(len(x['right']) for x in blocks.values()):2d}줄  {sum(i['qty'] for i in mine)}매")
    if unassigned:
        print(f"  ! 어느 묶음에도 안 들어간 줄 {len(unassigned)}개 (동/호: {sorted({(i['dong'], i['ho']) for i in unassigned})})")
    for w in warnings:
        print('  ! ' + w)
    print('저장:', out)


if __name__ == '__main__':
    main()
